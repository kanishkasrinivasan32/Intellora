import ast
import csv
import io
import json
from pathlib import Path

SUPPORTED = {'.txt', '.md', '.pdf', '.docx', '.csv', '.json', '.png', '.jpg', '.jpeg', '.webp', '.py', '.ipynb', '.sql'}

def read_file(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == '.pdf':
        import fitz
        with fitz.open(path) as doc:
            pages = [(i + 1, page.get_text().strip()) for i, page in enumerate(doc)]
            if not any(text for _, text in pages):
                raise ValueError('No readable text found. Scanned PDFs need OCR before upload.')
            return '\n\n'.join(f'## Page {number}\n{text}' for number, text in pages if text), {'pages': len(doc)}
    if ext == '.docx':
        from docx import Document
        doc = Document(path)
        rows = [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return '\n\n'.join([p.text for p in doc.paragraphs] + rows), {}
    if ext in {'.png', '.jpg', '.jpeg', '.webp'}:
        from PIL import Image
        import pytesseract
        try:
            with Image.open(path) as img: return pytesseract.image_to_string(img), {'reader': 'tesseract'}
        except pytesseract.TesseractNotFoundError:
            raise ValueError('Image OCR requires Tesseract. Install it and add it to PATH, then retry.')
    text = path.read_text(encoding='utf-8-sig', errors='replace')
    if ext == '.csv':
        rows = csv.DictReader(io.StringIO(text))
        return '\n\n'.join('; '.join(f'{k}: {v}' for k, v in row.items()) for row in rows), {}
    if ext == '.json':
        data = json.loads(text)
        def flatten(value, prefix=''):
            if isinstance(value, dict): return '\n'.join(flatten(v, f'{prefix}.{k}'.strip('.')) for k,v in value.items())
            if isinstance(value, list): return '\n'.join(flatten(v, f'{prefix}[{i}]') for i,v in enumerate(value))
            return f'{prefix}: {value}'
        return flatten(data), {}
    if ext == '.ipynb':
        notebook = json.loads(text)
        parts = []
        for cell in notebook.get('cells', []):
            parts.append(f'## {cell["cell_type"]}\n' + ''.join(cell.get('source', [])))
            for output in cell.get('outputs', []):
                parts.append(''.join(output.get('text', output.get('data', {}).get('text/plain', []))))
        return '\n\n'.join(parts), {}
    if ext == '.py':
        try:
            tree = ast.parse(text)
            docs = [ast.get_docstring(n) for n in ast.walk(tree) if isinstance(n, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))]
            return '\n\n'.join(['## Documentation']+[d for d in docs if d]+['## Source code', text]), {}
        except SyntaxError: return text, {'warning': 'Python syntax could not be parsed; source preserved.'}
    return text, {}
